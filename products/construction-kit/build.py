#!/usr/bin/env python3
"""Build the UK Construction Kit: CIS deductions tracker + domestic reverse charge (DRC) invoice (wave 3).

Output: dist/Construction_Kit_CIS_and_Reverse_Charge.xlsx
Facts checked against gov.uk 2026-09-25: CIS deduction 20% (registered), 30% (unregistered), 0% (gross payment
status), on the labour element only (not materials, VAT or CITB levy); monthly returns for tax months
(6th to 5th) due by the 19th, electronic payment by the 22nd; subcontractor statements within 14 days of the
tax month end. DRC wording follows HMRC VAT reverse charge for building and construction services guidance.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "dist" / "Construction_Kit_CIS_and_Reverse_Charge.xlsx"
HEAD, HEAD_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="7F3F00")
INPUT = PatternFill("solid", fgColor="FFF2CC")
AMBER = PatternFill("solid", fgColor="FFE699")
TITLE = Font(bold=True, size=14)
GBP = '£#,##0.00'
ROWS = 500
DISCLAIMER = "A record-keeping aid, not tax advice. Check gov.uk/what-is-the-construction-industry-scheme and HMRC's VAT reverse charge guidance."


def head(ws, labels, widths):
    for i, (label, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"


def start(wb):
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Construction Kit – CIS deductions tracker + reverse charge invoice", TITLE),
        (DISCLAIMER, Font(italic=True, color="C00000")),
        ("", None),
        ("For contractors (you pay subcontractors)", Font(bold=True)),
        ("1. Subcontractors: add each one with their verification status: Gross (0%), Registered (20%) or Unregistered (30%).", None),
        ("2. Payments: one row per payment. Split labour from materials; deduction is worked out on labour only.", None),
        ("3. Monthly return: totals per tax month (6th to 5th). Return due by the 19th, payment by the 22nd (electronic).", None),
        ("4. Statement: give each subcontractor a statement within 14 days of the end of the tax month.", None),
        ("", None),
        ("For subcontractors (deductions taken from you)", Font(bold=True)),
        ("5. Deductions suffered: log what contractors deducted; claim it against your tax bill in Self Assessment (or via your company's payroll).", None),
        ("", None),
        ("Reverse charge (VAT-registered subcontractors)", Font(bold=True)),
        ("6. DRC invoice: for CIS services to a VAT- and CIS-registered customer who isn't the end user, don't charge VAT; the customer accounts for it.", None),
    ]
    for r, (t, f) in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=t)
        c.alignment = Alignment(wrap_text=True)
        if f:
            c.font = f


def subcontractors(wb):
    ws = wb.create_sheet("Subcontractors")
    head(ws, ["Name", "UTR", "Verification number", "Status", "Rate"], [30, 16, 22, 16, 8])
    dv = DataValidation(type="list", formula1='"Gross,Registered,Unregistered"', allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(2, 102):
        for col in "ABCD":
            ws[f"{col}{r}"].fill = INPUT
        dv.add(f"D{r}")
        ws[f"E{r}"] = f'=IF(D{r}="Gross",0,IF(D{r}="Registered",0.2,IF(D{r}="Unregistered",0.3,"")))'
        ws[f"E{r}"].number_format = "0%"


def payments(wb):
    ws = wb.create_sheet("Payments")
    head(ws, ["Payment date", "Subcontractor", "Labour £", "Materials £", "VAT £", "Tax month ends", "Rate",
              "CIS deduction £", "Pay subcontractor £"], [13, 28, 12, 12, 10, 14, 7, 15, 17])
    dv = DataValidation(type="list", formula1="=Subcontractors!$A$2:$A$101", allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(2, ROWS + 2):
        for col in "ABCDE":
            ws[f"{col}{r}"].fill = INPUT
        dv.add(f"B{r}")
        ws[f"A{r}"].number_format = ws[f"F{r}"].number_format = "dd/mm/yyyy"
        # Tax month ends on the 5th: payments on the 6th onwards belong to the month ending next 5th.
        ws[f"F{r}"] = f'=IF(A{r}="","",IF(DAY(A{r})<=5,DATE(YEAR(A{r}),MONTH(A{r}),5),DATE(YEAR(A{r}),MONTH(A{r})+1,5)))'
        ws[f"G{r}"] = f'=IF(B{r}="","",IFERROR(INDEX(Subcontractors!$E$2:$E$101,MATCH(B{r},Subcontractors!$A$2:$A$101,0)),0.3))'
        ws[f"G{r}"].number_format = "0%"
        ws[f"H{r}"] = f'=IF(B{r}="","",ROUND(N(C{r})*G{r},2))'
        ws[f"I{r}"] = f'=IF(B{r}="","",N(C{r})+N(D{r})+N(E{r})-H{r})'
        for col in "CDEHI":
            ws[f"{col}{r}"].number_format = GBP
    ws["K1"] = "Unknown subcontractor → 30% until verified. Deduction is on labour only (not materials, VAT or CITB levy)."


def monthly(wb):
    ws = wb.create_sheet("Monthly return")
    head(ws, ["Tax month ends", "Return due", "Pay HMRC by (electronic)", "Labour paid £", "Materials £",
              "CIS deducted £", "Days to return"], [14, 13, 20, 14, 13, 14, 13])
    for i in range(12):
        r = i + 2
        ws[f"A{r}"] = f"=DATE(2026,{5 + i},5)"
        ws[f"B{r}"] = f"=A{r}+14"   # 19th
        ws[f"C{r}"] = f"=A{r}+17"   # 22nd
        ws[f"D{r}"] = f"=SUMIFS(Payments!$C:$C,Payments!$F:$F,A{r})"
        ws[f"E{r}"] = f"=SUMIFS(Payments!$D:$D,Payments!$F:$F,A{r})"
        ws[f"F{r}"] = f"=SUMIFS(Payments!$H:$H,Payments!$F:$F,A{r})"
        ws[f"G{r}"] = f"=B{r}-TODAY()"
        for col in "ABC":
            ws[f"{col}{r}"].number_format = "dd/mm/yyyy"
        for col in "DEF":
            ws[f"{col}{r}"].number_format = GBP
    ws.conditional_formatting.add("G2:G13", FormulaRule(formula=["AND(G2>=0,G2<=7)"], fill=AMBER))
    ws["I1"] = "File a return every month you're registered as a contractor, including a nil return if you paid nobody (or tell HMRC you're inactive)."


def statement(wb):
    ws = wb.create_sheet("Statement")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    ws["A1"], ws["A1"].font = "CIS payment and deduction statement", TITLE
    rows = [("Contractor name", None), ("Contractor employer's reference", None), ("Subcontractor", None),
            ("Subcontractor verification number", None), ("Tax month ending", None),
            ("Gross amount paid (excl. VAT) £", "=SUMIFS(Payments!$C:$C,Payments!$B:$B,B5,Payments!$F:$F,B7)+SUMIFS(Payments!$D:$D,Payments!$B:$B,B5,Payments!$F:$F,B7)"),
            ("Cost of materials £", "=SUMIFS(Payments!$D:$D,Payments!$B:$B,B5,Payments!$F:$F,B7)"),
            ("Amount liable to deduction £", "=B8-B9"),
            ("Amount deducted £", "=SUMIFS(Payments!$H:$H,Payments!$B:$B,B5,Payments!$F:$F,B7)")]
    for i, (label, val) in enumerate(rows, 3):
        ws[f"A{i}"] = label
        if val:
            ws[f"B{i}"] = val
            ws[f"B{i}"].number_format = GBP
        else:
            ws[f"B{i}"].fill = INPUT
    ws["B7"].number_format = "dd/mm/yyyy"
    ws["A13"] = "Give this to the subcontractor within 14 days of the end of the tax month. Keep a copy."


def suffered(wb):
    ws = wb.create_sheet("Deductions suffered")
    head(ws, ["Date", "Contractor", "Gross paid £", "Materials £", "CIS deducted £", "Statement received?"], [12, 28, 13, 13, 14, 18])
    for r in range(2, 202):
        for col in "ABCDEF":
            ws[f"{col}{r}"].fill = INPUT
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
        for col in "CDE":
            ws[f"{col}{r}"].number_format = GBP
    ws["H1"], ws["H1"].font = "Total deducted £", Font(bold=True)
    ws["I1"] = "=SUM(E2:E201)"
    ws["I1"].number_format = GBP
    ws["H2"] = "Enter this total in your Self Assessment ('CIS deductions') to reduce your tax bill or get a refund."


def drc_invoice(wb):
    ws = wb.create_sheet("DRC invoice")
    for col, w in zip("ABCDE", (40, 8, 12, 12, 16)):
        ws.column_dimensions[col].width = w
    ws["A1"], ws["A1"].font = "VAT INVOICE – DOMESTIC REVERSE CHARGE", TITLE
    for i, label in enumerate(["Your business name", "Your address", "Your VAT number", "Customer name", "Customer address",
                               "Customer VAT number", "Invoice number", "Invoice date / tax point"], 3):
        ws[f"A{i}"] = label
        ws[f"B{i}"].fill = INPUT
    for i, h in enumerate(["Description (mark reverse charge items)", "Qty", "Unit £", "VAT rate", "Net £"], 1):
        c = ws.cell(row=12, column=i, value=h)
        c.font, c.fill = HEAD, HEAD_FILL
    for r in range(13, 21):
        for col in "ABCD":
            ws[f"{col}{r}"].fill = INPUT
        ws[f"E{r}"] = f'=IF(B{r}="","",B{r}*C{r})'
        ws[f"E{r}"].number_format = GBP
        ws[f"C{r}"].number_format = GBP
        ws[f"F{r}"] = f'=IF(E{r}="",0,E{r}*N(D{r}))'  # hidden helper: VAT per line
    ws.column_dimensions["F"].hidden = True
    ws["D22"], ws["E22"] = "Total (net) £", "=SUM(E13:E20)"
    ws["D23"], ws["E23"] = "VAT to be accounted for by customer £", "=SUM(F13:F20)"
    ws["D24"], ws["E24"] = "Amount payable £", "=E22"
    for ref in ("E22", "E23", "E24"):
        ws[ref].number_format = GBP
    ws["A26"] = ("Reverse charge: Customer to account to HMRC for the reverse charge output tax on the VAT exclusive "
                 "price of items marked 'reverse charge'.")
    ws["A26"].font = Font(bold=True)
    ws["A26"].alignment = Alignment(wrap_text=True)
    ws["A28"] = ("Use only when the customer is VAT- and CIS-registered and is not an end user or intermediary. "
                 "Enter the VAT rate (e.g. 0.2) so the customer can see the VAT they account for. " + DISCLAIMER)
    ws["A28"].alignment = Alignment(wrap_text=True)


def build(out=OUT):
    wb = Workbook()
    start(wb)
    subcontractors(wb)
    payments(wb)
    monthly(wb)
    statement(wb)
    suffered(wb)
    drc_invoice(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build())
