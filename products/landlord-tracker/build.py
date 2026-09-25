#!/usr/bin/env python3
"""Build the UK (England) landlord compliance tracker.

Output: dist/Landlord_Compliance_Tracker_England.xlsx
Facts last verified against gov.uk / legislation.gov.uk: 2026-09-25.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "dist" / "Landlord_Compliance_Tracker_England.xlsx"
PROPERTIES = 10
HEAD, HEAD_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="375623")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
RED, AMBER, GREEN = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FFE699", "C6EFCE"))
TITLE = Font(bold=True, size=14)

DISCLAIMER = ("For landlords letting in England. A record-keeping aid, not legal advice. Rules differ in Wales, "
              "Scotland and Northern Ireland. Always check the linked gov.uk source; rules change.")

# (item, interval in months or None for one-off/event-based, rule summary, source)
CHECKS = [
    ("Gas safety check (CP12)", 12,
     "Every 12 months by a Gas Safe engineer; copy to tenants within 28 days; keep records 2 years.",
     "gov.uk/government/publications/gas-safety-landlords-and-tenants"),
    ("Electrical safety report (EICR)", 60,
     "At least every 5 years by a qualified person; copy to tenants within 28 days.",
     "gov.uk/government/publications/electrical-safety-standards-in-the-private-rented-sector-guidance-for-landlords-tenants-and-local-authorities"),
    ("Energy Performance Certificate (EPC)", 120,
     "Valid 10 years; rating must be E or above to let (unless exempt).",
     "gov.uk/guidance/domestic-private-rented-property-minimum-energy-efficiency-standard-landlord-guidance"),
    ("Smoke alarms tested", None,
     "At least one on every storey used as living accommodation; check working on the first day of each tenancy.",
     "gov.uk/government/publications/smoke-and-carbon-monoxide-alarms-explanatory-booklet-for-landlords"),
    ("Carbon monoxide alarms tested", None,
     "In any room with a fixed combustion appliance (not gas cookers); check on the first day of each tenancy.",
     "gov.uk/government/publications/smoke-and-carbon-monoxide-alarms-explanatory-booklet-for-landlords"),
    ("Deposit protected + prescribed information", None,
     "Protect in a government-approved scheme and give prescribed information within 30 days of receiving it.",
     "gov.uk/deposit-protection-schemes-and-landlords"),
    ("Right to Rent check", None,
     "Before the tenancy starts; repeat check before a time-limited permission expires.",
     "gov.uk/check-tenant-right-to-rent-documents"),
    ("Licence (HMO / selective) expiry", None,
     "If your council requires a licence, enter its expiry date. Licences usually last up to 5 years.",
     "gov.uk/house-in-multiple-occupation-licence"),
    ("Landlord insurance renewal", 12, "Your policy renewal date.", ""),
    ("Legionella risk assessment review", None,
     "Landlords must assess legionella risk. The law sets no fixed interval; review when anything changes.",
     "hse.gov.uk/legionnaires/landlords.htm"),
]

RRA = [
    ("1 May 2026", "Fixed-term assured shorthold tenancies became assured periodic tenancies. Section 21 'no fault' "
     "evictions ended; possession now only via Section 8 grounds.",
     "legislation.gov.uk/uksi/2026/421/made"),
    ("31 May 2026", "Existing written tenancies: give every named tenant the government's 'Renters' Rights Act "
     "Information Sheet 2026' (exact gov.uk PDF). Fine up to £7,000 if not given.",
     "gov.uk (search 'Renters' Rights Act Information Sheet 2026')"),
    ("New tenancies from 1 May 2026", "Provide a written statement of terms before the tenancy starts.",
     "gov.uk/guidance/renters-rights-act-2025"),
    ("Rent increases", "Only once a year, using the Section 13 notice procedure; tenants can challenge at tribunal.",
     "gov.uk/guidance/renters-rights-act-2025"),
    ("Pets", "Tenants can ask to keep a pet; you must not unreasonably refuse and must reply in writing.",
     "gov.uk/guidance/renters-rights-act-2025"),
    ("Discrimination", "Must not refuse tenants because they have children or receive benefits.",
     "gov.uk/guidance/renters-rights-act-2025"),
    ("Rental bidding", "Must advertise a rent and not accept offers above it.",
     "gov.uk/guidance/renters-rights-act-2025"),
    ("Later phases", "Private Rented Sector Database registration, Landlord Ombudsman membership and Decent Homes "
     "Standard - dates to be confirmed by government. Check before letting.",
     "gov.uk/guidance/renters-rights-act-2025"),
]


def head(ws, row, labels):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def start(wb):
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 110
    rows = [
        ("Landlord compliance tracker – England", TITLE),
        (DISCLAIMER, Font(italic=True, color="C00000")),
        ("", None),
        ("1. Properties: list up to 10 properties (yellow cells).", None),
        ("2. Compliance: for each property and check, enter the date it was last done (or the expiry date for licences).", None),
        ("   Next due and status fill in automatically: red = overdue, amber = due within 60 days, green = OK.", None),
        ("3. Dashboard: see every overdue and upcoming item at a glance.", None),
        ("4. Renters' Rights Act: checklist of changes now in force, with sources.", None),
        ("", None),
        ("Facts last checked 25 September 2026. Buyers get free updates for 12 months as rules change.", Font(italic=True)),
    ]
    for r, (t, f) in enumerate(rows, 1):
        c = ws.cell(row=r, column=1, value=t)
        c.alignment = Alignment(wrap_text=True)
        if f:
            c.font = f


def properties(wb):
    ws = wb.create_sheet("Properties")
    head(ws, 1, ["#", "Address", "Tenant name(s)", "Tenancy start", "Deposit scheme + ref", "Notes"])
    for col, w in zip("ABCDEF", (5, 40, 28, 14, 26, 30)):
        ws.column_dimensions[col].width = w
    for i in range(PROPERTIES):
        r = i + 2
        ws[f"A{r}"] = i + 1
        for col in "BCDEF":
            ws[f"{col}{r}"].fill = INPUT_FILL
        ws[f"D{r}"].number_format = "dd/mm/yyyy"


def compliance(wb):
    ws = wb.create_sheet("Compliance")
    head(ws, 1, ["Property #", "Address", "Check", "Last done / expiry", "Interval (months)", "Next due",
                 "Days left", "Status", "Rule", "Source"])
    for col, w in zip("ABCDEFGHIJ", (10, 32, 38, 16, 11, 14, 10, 11, 70, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D2"
    r = 2
    for p in range(1, PROPERTIES + 1):
        for item, months, rule, src in CHECKS:
            ws[f"A{r}"] = p
            ws[f"B{r}"] = f'=IF(Properties!B{p + 1}="","",Properties!B{p + 1})'
            ws[f"C{r}"] = item
            ws[f"D{r}"].fill = INPUT_FILL
            ws[f"D{r}"].number_format = "dd/mm/yyyy"
            ws[f"E{r}"] = months
            if months:
                ws[f"F{r}"] = f'=IF(D{r}="","",EDATE(D{r},E{r}))'
            elif item.startswith("Licence"):
                ws[f"F{r}"] = f'=IF(D{r}="","",D{r})'
            else:
                ws[f"F{r}"] = ""  # event-based: record date only
            ws[f"F{r}"].number_format = "dd/mm/yyyy"
            ws[f"G{r}"] = f'=IF(F{r}="","",F{r}-TODAY())'
            ws[f"H{r}"] = (f'=IF(B{r}="","",IF(F{r}="",IF(D{r}="","Missing","Recorded"),'
                           f'IF(G{r}<0,"OVERDUE",IF(G{r}<=60,"Due soon","OK"))))')
            ws[f"I{r}"] = rule
            ws[f"J{r}"] = src
            r += 1
    last = r - 1
    rng = f"H2:H{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['OR(H2="OVERDUE",H2="Missing")'], fill=RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['H2="Due soon"'], fill=AMBER))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['OR(H2="OK",H2="Recorded")'], fill=GREEN))
    ws.auto_filter.ref = f"A1:J{last}"
    return last


def dashboard(wb, last):
    ws = wb.create_sheet("Dashboard", 1)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    ws["A1"], ws["A1"].font = "Dashboard", TITLE
    for i, (label, crit) in enumerate([("Overdue", "OVERDUE"), ("Missing records", "Missing"),
                                       ("Due within 60 days", "Due soon"), ("OK", "OK")], 3):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = f'=COUNTIF(Compliance!H2:H{last},"{crit}")'
    ws.conditional_formatting.add("B3:B4", FormulaRule(formula=["B3>0"], fill=RED))
    ws.conditional_formatting.add("B5", FormulaRule(formula=["B5>0"], fill=AMBER))
    ws["A8"] = "Filter the Compliance sheet by Status to see which items need action."
    ws["A9"] = DISCLAIMER
    ws["A9"].font = Font(italic=True, color="C00000")


def rra(wb):
    ws = wb.create_sheet("Renters' Rights Act")
    head(ws, 1, ["When / topic", "What it means for you", "Done?", "Source"])
    for col, w in zip("ABCD", (26, 90, 8, 50)):
        ws.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1='"Yes,No,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    for r, (when, what, src) in enumerate(RRA, 2):
        ws[f"A{r}"], ws[f"B{r}"], ws[f"D{r}"] = when, what, src
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"C{r}"].fill = INPUT_FILL
        dv.add(f"C{r}")
    ws[f"A{len(RRA) + 3}"] = DISCLAIMER
    ws[f"A{len(RRA) + 3}"].font = Font(italic=True, color="C00000")


def build(out=OUT):
    wb = Workbook()
    start(wb)
    properties(wb)
    last = compliance(wb)
    dashboard(wb, last)
    rra(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build())
