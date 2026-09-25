#!/usr/bin/env python3
"""Build the "MTD Ready for Landlords" UK property income workbook (wave 2), reusing the MTD workbook engine.

Output: dist/MTD_Ready_Landlord_Workbook_2026-27.xlsx
Expense headings follow HMRC's UK property pages (SA105) used for MTD property updates. Residential finance
costs (mortgage interest) are kept separate: they are not deducted, they give a 20% basic-rate tax credit.
Furnished holiday lettings rules were abolished from 6 April 2025. Facts checked against gov.uk 2026-09-25.
"""
import importlib.util
from pathlib import Path

from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("mtd", ROOT / "products" / "mtd-workbook" / "build.py")
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)

OUT = ROOT / "dist" / "MTD_Ready_Landlord_Workbook_2026-27.xlsx"
m.INCOME_CATEGORIES = ["Rent received", "Premiums for the grant of a lease", "Reverse premiums and inducements", "Other property income"]
m.EXPENSE_CATEGORIES = [
    "Rent, rates, insurance and ground rents",
    "Property repairs and maintenance",
    "Non-residential property finance costs",
    "Legal, management and other professional fees",
    "Costs of services provided, including wages",
    "Travel costs",
    "Other allowable property expenses",
]
m.DISCLAIMER = (
    "This workbook helps landlords keep digital records. It is not HMRC-recognised software and cannot submit to "
    "HMRC by itself: use HMRC-recognised bridging software to send quarterly updates from a spreadsheet. "
    "It is not tax advice. Check gov.uk or an accountant for your circumstances."
)


def start_sheet(wb):
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("MTD Ready for Landlords – UK property income workbook, tax year 2026-27", m.TITLE),
        (m.DISCLAIMER, Font(italic=True, color="C00000")),
        ("", None),
        ("How to use", Font(bold=True)),
        ("1. Settings: choose Standard (6 Apr) or Calendar (1 Apr) quarters.", None),
        ("2. Income: one row per rent receipt or other property income. Jointly owned? Enter only your share.", None),
        ("3. Expenses: one row per cost; put any personal element in 'Private use £'.", None),
        ("4. Mortgage interest: record it on the 'Residential finance costs' sheet, not as an expense (see why there).", None),
        ("5. Quarterly summary shows the cumulative totals your bridging software asks for.", None),
        ("", None),
        ("Who must use MTD for Income Tax", Font(bold=True)),
        ("Gross self-employment + gross property income over £50,000 (2024-25 return) → from 6 April 2026; over £30,000 → 6 April 2027; over £20,000 → 6 April 2028.", None),
        ("Property allowance: if your gross property income is £1,000 or less you may not need to report it; above that you can deduct £1,000 instead of actual expenses.", None),
        ("Furnished holiday lettings rules ended on 6 April 2025; holiday lets are now treated like other property income.", None),
        ("", None),
        ("Facts checked against gov.uk on 25 September 2026. Buyers get free updates for 12 months.", Font(italic=True)),
    ]
    for r, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=text)
        c.alignment = Alignment(wrap_text=True)
        if font:
            c.font = font


def finance_sheet(wb):
    ws = wb.create_sheet("Residential finance costs")
    m.header(ws, 1, ["Date", "Lender", "Property", "Interest and finance costs £"], [12, 26, 30, 22])
    for r in range(2, 202):
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
        ws[f"D{r}"].number_format = m.GBP
    ws["F1"], ws["F1"].font = "Year total £", Font(bold=True)
    ws["G1"] = "=SUM(D2:D201)"
    ws["F2"], ws["G2"] = "Estimated 20% tax credit £", "=G1*0.2"
    ws["G1"].number_format = ws["G2"].number_format = m.GBP
    ws["F4"] = ("Mortgage interest on residential lets is not deducted from rental profit. Instead you get a tax "
                "credit of 20% of these costs (the credit can be limited; see gov.uk 'restricting finance cost relief').")
    ws["F4"].alignment = Alignment(wrap_text=True)
    ws.column_dimensions["F"].width = 60


def build(out=OUT):
    wb = m.Workbook()
    start_sheet(wb)
    m.settings_sheet(wb)
    m.entry_sheet(wb, "Income", ["Date", "Description", "Category", "Amount £", "Property / tenant"],
                  [12, 40, 34, 14, 24], "C", m.INCOME_CATEGORIES, ["D"])
    m.entry_sheet(wb, "Expenses", ["Date", "Supplier", "Description", "Category", "Amount £", "Private use £", "Business part £"],
                  [12, 24, 34, 46, 14, 14, 16], "D", m.EXPENSE_CATEGORIES, ["E", "F"], private=True)
    finance_sheet(wb)
    m.mileage_sheet(wb)
    m.home_sheet(wb)  # summary references it; hidden below because landlords rarely use it
    m.summary_sheet(wb)
    m.deadlines_sheet(wb)
    m.lists_sheet(wb)
    wb["Use of home"].sheet_state = "hidden"
    s = wb["Quarterly summary"]
    for row in s.iter_rows():
        c = row[0]
        if isinstance(c.value, str) and c.value.startswith("Mileage allowance"):
            c.value = "Mileage allowance (add to 'Travel costs')"
        elif isinstance(c.value, str) and c.value.startswith("Use of home"):
            for cell in row:
                cell.value = None
        elif c.value == "Profit before simplified expenses":
            c.value = "Rental profit before mileage and finance-cost credit"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build())
