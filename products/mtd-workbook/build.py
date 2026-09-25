#!/usr/bin/env python3
"""Build the "MTD Ready" sole-trader record-keeping workbook (2026-27 tax year).

Output: dist/MTD_Ready_Workbook_2026-27.xlsx  (opens in Excel and Google Sheets)
Facts last verified against gov.uk: 2026-09-25.
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "dist" / "MTD_Ready_Workbook_2026-27.xlsx"
ROWS = 1000  # data rows available on each entry sheet

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE = Font(bold=True, size=14)
GBP = '£#,##0.00'

INCOME_CATEGORIES = ["Turnover", "Other business income"]
# Expense categories used by HMRC's MTD self-employment quarterly update (match SA103F headings).
EXPENSE_CATEGORIES = [
    "Cost of goods bought for resale or goods used",
    "Construction industry - payments to subcontractors",
    "Wages, salaries and other staff costs",
    "Car, van and travel expenses",
    "Rent, rates, power and insurance costs",
    "Repairs and maintenance of property and equipment",
    "Phone, fax, stationery and other office costs",
    "Advertising costs",
    "Business entertainment costs",
    "Interest on bank and other loans",
    "Bank, credit card and other financial charges",
    "Irrecoverable debts written off",
    "Accountancy, legal and other professional fees",
    "Depreciation and loss or profit on sale of assets",
    "Other business expenses",
]

DISCLAIMER = (
    "This workbook helps you keep digital records. It is not HMRC-recognised software and cannot submit "
    "to HMRC by itself: use HMRC-recognised bridging software to send quarterly updates from a spreadsheet. "
    "It is not tax advice. Check gov.uk or an accountant for your circumstances."
)


def header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def start_sheet(wb):
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("MTD Ready – sole-trader record-keeping workbook, tax year 2026-27", TITLE),
        (DISCLAIMER, Font(italic=True, color="C00000")),
        ("", None),
        ("How to use", Font(bold=True)),
        ("1. Settings: enter your business name and choose Standard (6 Apr) or Calendar (1 Apr) quarters.", None),
        ("2. Income and Expenses: add one row per transaction. Yellow cells are for you; pick categories from the drop-down.", None),
        ("3. If part of an expense is personal, put that part in 'Private use £' – only the business part is counted.", None),
        ("4. Mileage and Use of home work out HMRC flat rates if you use simplified expenses.", None),
        ("5. Quarterly summary shows each quarter and the running year-to-date totals your software asks for.", None),
        ("6. Deadlines lists every MTD date for 2026-27.", None),
        ("", None),
        ("Who must use MTD for Income Tax", Font(bold=True)),
        ("From 6 April 2026: qualifying self-employment + property income over £50,000 (based on 2024-25 return).", None),
        ("From 6 April 2027: over £30,000.   From 6 April 2028: over £20,000.", None),
        ("Source: gov.uk/guidance/check-if-youre-eligible-for-making-tax-digital-for-income-tax", None),
        ("", None),
        ("Facts last checked against gov.uk on 25 September 2026. Rules can change – buyers get free updates for 12 months.", Font(italic=True)),
    ]
    for r, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=text)
        c.alignment = Alignment(wrap_text=True)
        if font:
            c.font = font


def settings_sheet(wb):
    ws = wb.create_sheet("Settings")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 18
    ws["A1"], ws["A1"].font = "Settings", TITLE
    ws["A3"], ws["B3"] = "Business name", ""
    ws["A4"], ws["B4"] = "Quarter basis (Standard / Calendar)", "Standard"
    for ref in ("B3", "B4"):
        ws[ref].fill = INPUT_FILL
    dv = DataValidation(type="list", formula1='"Standard,Calendar"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")
    ws["A6"], ws["B6"], ws["C6"] = "Quarter", "Start", "End"
    for c in ("A6", "B6", "C6"):
        ws[c].font, ws[c].fill = HEAD, HEAD_FILL
    std = [(date(2026, 4, 6), date(2026, 7, 5)), (date(2026, 7, 6), date(2026, 10, 5)),
           (date(2026, 10, 6), date(2027, 1, 5)), (date(2027, 1, 6), date(2027, 4, 5))]
    cal = [(date(2026, 4, 1), date(2026, 6, 30)), (date(2026, 7, 1), date(2026, 9, 30)),
           (date(2026, 10, 1), date(2026, 12, 31)), (date(2027, 1, 1), date(2027, 3, 31))]
    for i, ((s1, e1), (s2, e2)) in enumerate(zip(std, cal)):
        r = 7 + i
        ws.cell(row=r, column=1, value=f"Q{i + 1}")
        for col, a, b in ((2, s1, s2), (3, e1, e2)):
            c = ws.cell(row=r, column=col,
                        value=f'=IF($B$4="Calendar",DATE({b.year},{b.month},{b.day}),DATE({a.year},{a.month},{a.day}))')
            c.number_format = "dd/mm/yyyy"


def entry_sheet(wb, title, cols, widths, category_col, categories, amount_cols, private=False):
    ws = wb.create_sheet(title)
    header(ws, 1, cols, widths)
    dv = DataValidation(type="list", formula1=f"=Lists!${'A' if title == 'Income' else 'B'}$2:${'A' if title == 'Income' else 'B'}${len(categories) + 1}",
                        allow_blank=True, showErrorMessage=True, error="Pick a category from the list")
    ws.add_data_validation(dv)
    dv.add(f"{category_col}2:{category_col}{ROWS + 1}")
    for r in range(2, ROWS + 2):
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        for col in amount_cols:
            ws[f"{col}{r}"].number_format = GBP
        if private:
            ws[f"G{r}"] = f'=IF(E{r}="","",E{r}-N(F{r}))'
            ws[f"G{r}"].number_format = GBP
    return ws


def mileage_sheet(wb):
    ws = wb.create_sheet("Mileage")
    header(ws, 1, ["Date", "From", "To", "Business purpose", "Miles", "Miles so far this year", "Allowance £"],
           [12, 18, 18, 34, 10, 14, 14])
    for r in range(2, ROWS + 2):
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
        prev = "0" if r == 2 else f"F{r - 1}"
        ws[f"F{r}"] = f'=IF(E{r}="",{prev},{prev}+E{r})'
        # 45p for the first 10,000 business miles in the year, 25p after (cars and vans)
        ws[f"G{r}"] = (f'=IF(E{r}="","",MIN(E{r},MAX(0,10000-{prev}))*0.45'
                       f'+MAX(0,E{r}-MAX(0,10000-{prev}))*0.25)')
        ws[f"G{r}"].number_format = GBP
    ws["I1"] = "Cars and vans: 45p/mile first 10,000 business miles, then 25p. Motorcycles: 24p (not calculated here)."
    ws["I2"] = "You cannot claim mileage and actual running costs for the same vehicle. Source: gov.uk/simpler-income-tax-simplified-expenses"


def home_sheet(wb):
    ws = wb.create_sheet("Use of home")
    header(ws, 1, ["Month", "Hours worked at home", "Flat rate £"], [14, 22, 14])
    months = ["Apr 2026", "May 2026", "Jun 2026", "Jul 2026", "Aug 2026", "Sep 2026", "Oct 2026",
              "Nov 2026", "Dec 2026", "Jan 2027", "Feb 2027", "Mar 2027"]
    for i, m in enumerate(months, 2):
        ws[f"A{i}"] = m
        ws[f"B{i}"].fill = INPUT_FILL
        ws[f"C{i}"] = f'=IF(B{i}>=101,26,IF(B{i}>=51,18,IF(B{i}>=25,10,0)))'
        ws[f"C{i}"].number_format = GBP
    ws["A14"], ws["C14"] = "Total", "=SUM(C2:C13)"
    ws["A14"].font = Font(bold=True)
    ws["C14"].number_format = GBP
    ws["E1"] = "HMRC flat rates: 25-50 hours £10/month, 51-100 £18, 101+ £26. Source: gov.uk/simpler-income-tax-simplified-expenses"


def summary_sheet(wb):
    ws = wb.create_sheet("Quarterly summary")
    ws.column_dimensions["A"].width = 52
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 14
    ws["A1"], ws["A1"].font = "Quarterly summary – copy these into your bridging software", TITLE
    ws["A2"] = "Quarterly updates are cumulative: HMRC asks for totals from the start of the tax year."
    labels = ["Category", "Q1", "Q2", "Q3", "Q4", "YTD to Q1", "YTD to Q2", "YTD to Q3", "YTD to Q4"]
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=4, column=i, value=label)
        c.font, c.fill = HEAD, HEAD_FILL

    def row(r, label, sheet, amt, cat, cat_value):
        ws.cell(row=r, column=1, value=label)
        for q in range(4):
            s, e = f"Settings!$B${7 + q}", f"Settings!$C${7 + q}"
            crit = f'{sheet}!${cat}:${cat},"{cat_value}",' if cat else ""
            ws.cell(row=r, column=2 + q,
                    value=f'=SUMIFS({sheet}!${amt}:${amt},{crit}{sheet}!$A:$A,">="&{s},{sheet}!$A:$A,"<="&{e})')
            ytd_s = "Settings!$B$7"
            ws.cell(row=r, column=6 + q,
                    value=f'=SUMIFS({sheet}!${amt}:${amt},{crit}{sheet}!$A:$A,">="&{ytd_s},{sheet}!$A:$A,"<="&{e})')
        for col in range(2, 10):
            ws.cell(row=r, column=col).number_format = GBP

    r = 5
    ws.cell(row=r, column=1, value="INCOME").font = Font(bold=True)
    r += 1
    first_income = r
    for cat in INCOME_CATEGORIES:
        row(r, cat, "Income", "D", "C", cat)
        r += 1
    last_income = r - 1
    r += 1
    ws.cell(row=r, column=1, value="EXPENSES (business part only)").font = Font(bold=True)
    r += 1
    first_exp = r
    for cat in EXPENSE_CATEGORIES:
        row(r, cat, "Expenses", "G", "D", cat)
        r += 1
    last_exp = r - 1
    r += 1
    for label, fn in (("Total income", f"SUM({{c}}{first_income}:{{c}}{last_income})"),
                      ("Total expenses", f"SUM({{c}}{first_exp}:{{c}}{last_exp})")):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        for col in "BCDEFGHI":
            c = ws[f"{col}{r}"]
            c.value, c.number_format, c.font = "=" + fn.format(c=col), GBP, Font(bold=True)
        r += 1
    ws.cell(row=r, column=1, value="Profit before simplified expenses").font = Font(bold=True)
    for col in "BCDEFGHI":
        c = ws[f"{col}{r}"]
        c.value, c.number_format, c.font = f"={col}{r - 2}-{col}{r - 1}", GBP, Font(bold=True)
    r += 2
    ws.cell(row=r, column=1, value="Simplified expenses (add to the category your software asks for)").font = Font(bold=True)
    r += 1
    row(r, "Mileage allowance (Car, van and travel)", "Mileage", "G", None, None)
    r += 2
    ws.cell(row=r, column=1, value="Use of home flat rate, full year (see Use of home sheet)")
    ws.cell(row=r, column=9, value="='Use of home'!C14").number_format = GBP
    r += 2
    ws.cell(row=r, column=1, value=DISCLAIMER).font = Font(italic=True, color="C00000")


def deadlines_sheet(wb):
    ws = wb.create_sheet("Deadlines")
    header(ws, 1, ["What", "Covers", "Deadline", "Days left"], [40, 34, 14, 12])
    items = [
        ("Quarterly update 1", "6 Apr - 5 Jul 2026 (cumulative)", date(2026, 8, 7)),
        ("Quarterly update 2", "Year to 5 Oct 2026", date(2026, 11, 7)),
        ("Quarterly update 3", "Year to 5 Jan 2027", date(2027, 2, 7)),
        ("Quarterly update 4", "Year to 5 Apr 2027", date(2027, 5, 7)),
        ("Final declaration + pay balance", "Tax year 2026-27", date(2028, 1, 31)),
    ]
    for i, (what, cover, d) in enumerate(items, 2):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = what, cover, d
        ws[f"C{i}"].number_format = "dd/mm/yyyy"
        ws[f"D{i}"] = f'=C{i}-TODAY()'
    ws.conditional_formatting.add("D2:D6", FormulaRule(formula=["AND(D2>=0,D2<=14)"],
                                                       fill=PatternFill("solid", fgColor="F8CBAD")))
    ws["A8"] = "Same deadlines apply whether you use standard or calendar quarters. Source: gov.uk (MTD for Income Tax)."


def lists_sheet(wb):
    ws = wb.create_sheet("Lists")
    ws["A1"], ws["B1"] = "Income categories", "Expense categories"
    for i, c in enumerate(INCOME_CATEGORIES, 2):
        ws[f"A{i}"] = c
    for i, c in enumerate(EXPENSE_CATEGORIES, 2):
        ws[f"B{i}"] = c
    ws.sheet_state = "hidden"


def build(out=OUT):
    wb = Workbook()
    start_sheet(wb)
    settings_sheet(wb)
    entry_sheet(wb, "Income", ["Date", "Description", "Category", "Amount £", "Invoice / ref"],
                [12, 40, 26, 14, 18], "C", INCOME_CATEGORIES, ["D"])
    entry_sheet(wb, "Expenses", ["Date", "Supplier", "Description", "Category", "Amount £", "Private use £", "Business part £"],
                [12, 24, 34, 46, 14, 14, 16], "D", EXPENSE_CATEGORIES, ["E", "F"], private=True)
    mileage_sheet(wb)
    home_sheet(wb)
    summary_sheet(wb)
    deadlines_sheet(wb)
    lists_sheet(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build())
