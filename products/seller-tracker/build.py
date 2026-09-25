#!/usr/bin/env python3
"""Build the Online Seller & Side-Hustle Tracker (wave 3: ideas W058 + W088).

Output: dist/Online_Seller_and_Side_Hustle_Tracker.xlsx
Facts checked 2026-09-25: platforms (eBay, Vinted, Etsy, Airbnb...) report sellers to HMRC who make 30+ sales or
more than €2,000 (about £1,700) in a calendar year. Reporting is not tax: selling your own used belongings is
usually not trading. Trading income above the £1,000 trading allowance (per tax year, gross) must be reported.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "dist" / "Online_Seller_and_Side_Hustle_Tracker.xlsx"
HEAD, HEAD_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4A3AA7")
INPUT = PatternFill("solid", fgColor="FFF2CC")
RED, AMBER, GREEN = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FFE699", "C6EFCE"))
TITLE = Font(bold=True, size=14)
GBP = '£#,##0.00'
PLATFORMS = ["eBay", "Vinted", "Etsy", "Depop", "Amazon", "Facebook Marketplace", "Airbnb", "Other"]
DISCLAIMER = ("A record-keeping aid, not tax advice. See gov.uk 'Selling online and paying tax' and "
              "'Tax-free allowances on property and trading income'.")


def start(wb):
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 110
    for r, (t, f) in enumerate([
        ("Online Seller & Side-Hustle Tracker", TITLE),
        (DISCLAIMER, Font(italic=True, color="C00000")),
        ("", None),
        ("Two different things people mix up:", Font(bold=True)),
        ("1. REPORTING: platforms send HMRC your details if you make 30+ sales or over €2,000 (about £1,700) in a calendar year. That alone is not tax.", None),
        ("2. TAX: you pay tax only on TRADING (buying or making things to sell for profit), when gross trading income is over the £1,000 trading allowance in a tax year (6 April – 5 April).", None),
        ("Selling your own used clothes and belongings for less than you paid is normally not trading, however many you sell.", None),
        ("", None),
        ("How to use", Font(bold=True)),
        ("Sales: one row per sale. Mark 'Type' as Personal item or Trading. The Dashboard shows both tests for each platform and year.", None),
        ("Badges of trade: answer the questions if you're unsure whether you're trading.", None),
    ], 1):
        c = ws.cell(row=r, column=1, value=t)
        c.alignment = Alignment(wrap_text=True)
        if f:
            c.font = f


def sales(wb):
    ws = wb.create_sheet("Sales")
    labels = ["Date", "Platform", "Item", "Type", "Sale price £", "Postage charged £", "What it cost you £",
              "Platform fees £", "Postage paid £", "Profit £", "Tax year"]
    widths = [12, 20, 30, 15, 12, 12, 14, 12, 12, 11, 10]
    for i, (l, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row=1, column=i, value=l)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    pdv = DataValidation(type="list", formula1='"' + ",".join(PLATFORMS) + '"', allow_blank=True)
    tdv = DataValidation(type="list", formula1='"Personal item,Trading"', allow_blank=True)
    ws.add_data_validation(pdv)
    ws.add_data_validation(tdv)
    for r in range(2, 1002):
        for col in "ABCDEFGHI":
            ws[f"{col}{r}"].fill = INPUT
        pdv.add(f"B{r}")
        tdv.add(f"D{r}")
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
        ws[f"J{r}"] = f'=IF(A{r}="","",N(E{r})+N(F{r})-N(G{r})-N(H{r})-N(I{r}))'
        # Tax year label, e.g. 2026-27 for 6 Apr 2026 – 5 Apr 2027
        ws[f"K{r}"] = (f'=IF(A{r}="","",IF(A{r}>=DATE(YEAR(A{r}),4,6),YEAR(A{r})&"-"&RIGHT(YEAR(A{r})+1,2),'
                       f'(YEAR(A{r})-1)&"-"&RIGHT(YEAR(A{r}),2)))')
        for col in "EFGHIJ":
            ws[f"{col}{r}"].number_format = GBP


def dashboard(wb):
    ws = wb.create_sheet("Dashboard", 1)
    ws.column_dimensions["A"].width = 24
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
    ws["A1"], ws["A1"].font = "Dashboard", TITLE
    ws["A2"], ws["B2"] = "Calendar year for platform reporting", 2026
    ws["A3"], ws["B3"] = "Tax year for the trading allowance", "2026-27"
    ws["A4"], ws["B4"] = "€2,000 in £ (approx.; edit if the rate moves)", 1700
    for ref in ("B2", "B3", "B4"):
        ws[ref].fill = INPUT
    ws["B4"].number_format = GBP
    for i, h in enumerate(["Platform", "Sales (calendar yr)", "Takings £ (calendar yr)", "Likely reported?"], 1):
        c = ws.cell(row=6, column=i, value=h)
        c.font, c.fill = HEAD, HEAD_FILL
    for k, p in enumerate(PLATFORMS):
        r = 7 + k
        ws[f"A{r}"] = p
        crit = f'Sales!$B:$B,A{r},Sales!$A:$A,">="&DATE($B$2,1,1),Sales!$A:$A,"<="&DATE($B$2,12,31)'
        ws[f"B{r}"] = f"=COUNTIFS({crit})"
        ws[f"C{r}"] = f"=SUMIFS(Sales!$E:$E,{crit})"
        ws[f"C{r}"].number_format = GBP
        ws[f"D{r}"] = f'=IF(OR(B{r}>=30,C{r}>$B$4),"Yes – expect a report",IF(OR(B{r}>=24,C{r}>$B$4*0.8),"Close","No"))'
    last = 6 + len(PLATFORMS)
    ws.conditional_formatting.add(f"D7:D{last}", FormulaRule(formula=['LEFT(D7,3)="Yes"'], fill=AMBER))
    r = last + 2
    ws[f"A{r}"], ws[f"A{r}"].font = "Trading allowance test (tax year)", Font(bold=True)
    ws[f"A{r + 1}"] = "Gross trading income £"
    ws[f"B{r + 1}"] = f'=SUMIFS(Sales!$E:$E,Sales!$D:$D,"Trading",Sales!$K:$K,$B$3)+SUMIFS(Sales!$F:$F,Sales!$D:$D,"Trading",Sales!$K:$K,$B$3)'
    ws[f"A{r + 2}"] = "Trading profit £"
    ws[f"B{r + 2}"] = f'=SUMIFS(Sales!$J:$J,Sales!$D:$D,"Trading",Sales!$K:$K,$B$3)'
    ws[f"A{r + 3}"] = "What to do"
    ws[f"B{r + 3}"] = (f'=IF(B{r + 1}<=1000,"Under £1,000: covered by the trading allowance, nothing to report",'
                       f'"Over £1,000: register for Self Assessment; deduct £1,000 OR actual costs, whichever is better")')
    for rr in (r + 1, r + 2):
        ws[f"B{rr}"].number_format = GBP
    ws.conditional_formatting.add(f"B{r + 3}", FormulaRule(formula=[f'LEFT(B{r + 3},4)="Over"'], fill=RED))
    ws.conditional_formatting.add(f"B{r + 3}", FormulaRule(formula=[f'LEFT(B{r + 3},5)="Under"'], fill=GREEN))
    ws[f"A{r + 5}"] = DISCLAIMER
    ws[f"A{r + 5}"].font = Font(italic=True, color="C00000")


def badges(wb):
    ws = wb.create_sheet("Badges of trade")
    ws.column_dimensions["A"].width = 80
    ws.column_dimensions["B"].width = 12
    ws["A1"], ws["A1"].font = "Am I trading? HMRC's 'badges of trade' (answer Yes/No)", TITLE
    qs = ["Do you buy things with the intention of selling them at a profit?",
          "Do you make things to sell?",
          "Do you sell the same kind of item repeatedly?",
          "Do you change or repair items to sell them for more?",
          "Do you sell in an organised way (stock, listings, a shop name)?",
          "Did you own the items only briefly before selling?",
          "Is selling a regular source of income for you?"]
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, q in enumerate(qs, 3):
        ws[f"A{i}"] = q
        ws[f"B{i}"].fill = INPUT
        dv.add(f"B{i}")
    n = 3 + len(qs)
    ws[f"A{n + 1}"] = "Result"
    ws[f"B{n + 1}"] = f'=IF(COUNTIF(B3:B{n - 1},"Yes")>=2,"Likely trading",IF(COUNTIF(B3:B{n - 1},"Yes")=1,"Possibly","Unlikely"))'
    ws[f"A{n + 2}"] = "No single answer decides it. If you're unsure, check gov.uk or ask HMRC. " + DISCLAIMER
    ws[f"A{n + 2}"].alignment = Alignment(wrap_text=True)


def build(out=OUT):
    wb = Workbook()
    start(wb)
    sales(wb)
    dashboard(wb)
    badges(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build())
