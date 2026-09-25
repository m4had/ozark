#!/usr/bin/env python3
"""Build the UK small-business 12-month cash-flow forecast + VAT toolkit (idea #4).

Output: dist/Cashflow_and_VAT_Toolkit_UK.xlsx
Facts last verified 2026-09-25: VAT registration threshold £90,000 (rolling 12 months); Flat Rate Scheme joining
limit £150,000; limited cost trader rate 16.5% (goods < 2% of VAT-inclusive turnover or < £1,000 a year).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parents[2] / "dist" / "Cashflow_and_VAT_Toolkit_UK.xlsx"
HEAD, HEAD_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")
RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFE699")
TITLE = Font(bold=True, size=14)
GBP = '£#,##0;[Red]-£#,##0'
MONTHS = 12
COLS = [chr(66 + i) for i in range(MONTHS)]  # B..M
RECEIPTS = ["Sales (cash/card)", "Invoices paid by customers", "Other income", "Loans / owner money in"]
PAYMENTS = ["Stock / materials", "Wages", "Rent & rates", "Utilities", "Phone & internet", "Insurance",
            "Marketing", "Software & subscriptions", "Vehicle & travel", "Loan repayments", "VAT to HMRC",
            "Income tax / NI set-aside", "Owner drawings", "Other"]
DISCLAIMER = "A planning aid, not tax advice. Check gov.uk/vat-registration and gov.uk/vat-flat-rate-scheme."


def cashflow(wb):
    ws = wb.active
    ws.title = "Cash-flow forecast"
    ws.column_dimensions["A"].width = 30
    for c in COLS:
        ws.column_dimensions[c].width = 11
    ws["A1"], ws["A1"].font = "12-month cash-flow forecast", TITLE
    ws["A2"], ws["B2"] = "First month (date)", "=DATE(2026,10,1)"
    ws["B2"].fill, ws["B2"].number_format = INPUT, "mmm yyyy"
    ws["A3"], ws["B3"] = "Opening bank balance £", 0
    ws["B3"].fill, ws["B3"].number_format = INPUT, GBP
    ws["A5"] = "Month"
    for i, c in enumerate(COLS):
        ws[f"{c}5"] = f"=EDATE($B$2,{i})"
        ws[f"{c}5"].number_format = "mmm yy"
    for cell in ws[5]:
        cell.font, cell.fill = HEAD, HEAD_FILL
    r = 6
    ws.cell(row=r, column=1, value="MONEY IN").font = Font(bold=True)
    r += 1
    first_in = r
    for label in RECEIPTS:
        ws.cell(row=r, column=1, value=label)
        for c in COLS:
            ws[f"{c}{r}"].fill, ws[f"{c}{r}"].number_format = INPUT, GBP
        r += 1
    total_in = r
    ws.cell(row=r, column=1, value="Total in").font = Font(bold=True)
    for c in COLS:
        ws[f"{c}{r}"] = f"=SUM({c}{first_in}:{c}{r - 1})"
    r += 2
    ws.cell(row=r, column=1, value="MONEY OUT").font = Font(bold=True)
    r += 1
    first_out = r
    for label in PAYMENTS:
        ws.cell(row=r, column=1, value=label)
        for c in COLS:
            ws[f"{c}{r}"].fill, ws[f"{c}{r}"].number_format = INPUT, GBP
        r += 1
    total_out = r
    ws.cell(row=r, column=1, value="Total out").font = Font(bold=True)
    for c in COLS:
        ws[f"{c}{r}"] = f"=SUM({c}{first_out}:{c}{r - 1})"
    r += 2
    net, opening, closing = r, r + 1, r + 2
    ws.cell(row=net, column=1, value="Net cash flow").font = Font(bold=True)
    ws.cell(row=opening, column=1, value="Opening balance")
    ws.cell(row=closing, column=1, value="Closing balance").font = Font(bold=True)
    for i, c in enumerate(COLS):
        ws[f"{c}{net}"] = f"={c}{total_in}-{c}{total_out}"
        ws[f"{c}{opening}"] = "=$B$3" if i == 0 else f"={COLS[i - 1]}{closing}"
        ws[f"{c}{closing}"] = f"={c}{opening}+{c}{net}"
    for row in range(first_in, closing + 1):
        for c in COLS:
            ws[f"{c}{row}"].number_format = GBP
    ws.conditional_formatting.add(f"B{closing}:M{closing}", FormulaRule(formula=[f"B{closing}<0"], fill=RED))
    ws.cell(row=closing + 2, column=1, value="Red closing balance = you would be overdrawn that month.")
    ws.freeze_panes = "B6"


def vat_threshold(wb):
    ws = wb.create_sheet("VAT threshold watch")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 16
    ws["A1"], ws["A1"].font = "VAT registration threshold watch (rolling 12 months)", TITLE
    ws["A2"] = "You must register if VAT-taxable turnover in the last 12 months goes over £90,000 (or you expect to in the next 30 days alone)."
    ws["A3"], ws["B3"] = "Threshold £", 90000
    ws["B3"].number_format = GBP
    for i, h in enumerate(["Month", "Taxable turnover £", "Rolling 12-month total £", "Headroom £"], 1):
        cell = ws.cell(row=5, column=i, value=h)
        cell.font, cell.fill = HEAD, HEAD_FILL
    for i in range(24):
        r = 6 + i
        ws[f"A{r}"] = f"=EDATE(DATE(2025,4,1),{i})"
        ws[f"A{r}"].number_format = "mmm yyyy"
        ws[f"B{r}"].fill, ws[f"B{r}"].number_format = INPUT, GBP
        start = max(6, r - 11)
        ws[f"C{r}"] = f"=SUM(B{start}:B{r})"
        ws[f"D{r}"] = f"=$B$3-C{r}"
        ws[f"C{r}"].number_format = ws[f"D{r}"].number_format = GBP
    ws.conditional_formatting.add("C6:C29", FormulaRule(formula=["C6>$B$3"], fill=RED))
    ws.conditional_formatting.add("C6:C29", FormulaRule(formula=["C6>$B$3*0.85"], fill=AMBER))
    ws["F5"] = "Red = over the threshold: register within 30 days of the end of that month. Amber = within 15%."
    ws["F6"] = DISCLAIMER


def flat_rate(wb):
    ws = wb.create_sheet("Flat Rate vs Standard")
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 16
    ws["A1"], ws["A1"].font = "Flat Rate Scheme vs standard VAT accounting (annual)", TITLE
    rows = [
        ("Turnover excluding VAT £", 60000, True),
        ("VAT rate charged to customers", 0.2, True),
        ("Your sector flat rate % (from gov.uk list)", 0.145, True),
        ("VAT-registered less than 12 months? (Yes/No)", "No", True),
        ("Spend on GOODS incl. VAT £ (not services, fuel, food for staff, capital items)", 800, True),
        ("VAT you can reclaim on purchases under standard accounting £", 1500, True),
        ("", None, False),
        ("VAT-inclusive turnover £", "=B3*(1+B4)", False),
        ("Limited cost trader? (goods < 2% of VAT-incl. turnover or < £1,000)", '=IF(B7<MAX(B10*0.02,1000),"Yes","No")', False),
        ("Flat rate that applies", '=IF(B11="Yes",0.165,B5)-IF(B6="Yes",0.01,0)', False),
        ("VAT you pay HMRC on Flat Rate Scheme £", "=B10*B12", False),
        ("VAT you pay HMRC on standard accounting £", "=B3*B4-B8", False),
        ("Flat Rate Scheme saves you £ (negative = costs you)", "=B14-B13", False),
        ("Can you join? (turnover excl. VAT expected ≤ £150,000)", '=IF(B3<=150000,"Yes","No")', False),
    ]
    for i, (label, val, is_input) in enumerate(rows, 3):
        ws[f"A{i}"] = label
        if val is not None:
            ws[f"B{i}"] = val
        if is_input:
            ws[f"B{i}"].fill = INPUT
    for ref in ("B3", "B7", "B8", "B10", "B13", "B14", "B15"):
        ws[ref].number_format = GBP
    for ref in ("B4", "B5", "B12"):
        ws[ref].number_format = "0.0%"
    ws["A15"].font = ws["B15"].font = Font(bold=True)
    ws["A18"] = "Sector rates are listed at gov.uk/vat-flat-rate-scheme/how-much-you-pay. " + DISCLAIMER
    ws["A18"].alignment = Alignment(wrap_text=True)


def build(out=OUT):
    wb = Workbook()
    cashflow(wb)
    vat_threshold(wb)
    flat_rate(wb)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print(build())
